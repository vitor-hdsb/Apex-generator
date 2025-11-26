import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from PIL import Image, ImageTk
import pandas as pd
from fpdf import FPDF
from datetime import datetime
import os

class PDF(FPDF):
    def header(self):
        self.set_font('Arial', 'B', 14)
        self.cell(0, 10, 'TERMO DE CONSENTIMENTO', ln=True, align='C')
        self.ln(10)

    def footer(self):
        pass

def gerar_pdf_termos(resultados, df1, caminho_saida_pdf):
    pdf = PDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    
    for _, row in resultados.iterrows():
        nome = f"{row['First Name*']} {row['Last Name*']}"
        employee_id = row['Employee ID*']
        data_atual = datetime.now().strftime('%d/%m/%Y')

        login_info = df1.loc[df1['Employee ID'] == employee_id, 'Login']
        if not login_info.empty:
            login = login_info.values[0]
        else:
            login = "NÃO ENCONTRADO"

        texto_termo = (
            f"TERMO DE CONSENTIMENTO PARA UTILIZAÇÃO DE SISTEMAS INTERNOS\n\n"
            f"Eu, {nome}, portador do login {login} e Employee ID {employee_id}, "
            f"declaro que recebi treinamento sobre a utilização dos sistemas internos da empresa, "
            f"estando ciente das políticas de segurança da informação, responsabilidade sobre o acesso, "
            f"sigilo de dados e boas práticas de utilização.\n\n"
            f"Comprometo-me a utilizar os sistemas de forma ética, responsável e de acordo com as normas "
            f"estabelecidas pela companhia, ciente de que qualquer desvio poderá acarretar medidas disciplinares "
            f"e legais cabíveis.\n\n"
            f"Por ser verdade, firmo o presente termo para que produza seus efeitos legais.\n\n"
            f"Data: {data_atual}\n\n"
            f"Assinatura: ________________________________________________"
        )

        pdf.add_page()
        pdf.set_font('Arial', '', 12)
        pdf.multi_cell(0, 10, texto_termo)

    pdf.output(caminho_saida_pdf)

def salvar_excel_sem_formatacao(df, caminho):
    """Salva Excel sem qualquer formatação"""
    from openpyxl.styles import Font, PatternFill, Border, Alignment
    
    # Converter todos os dados para string
    df_clean = df.astype(str)
    df_clean = df_clean.replace('nan', '')
    
    with pd.ExcelWriter(caminho, engine='openpyxl') as writer:
        df_clean.to_excel(writer, index=False, sheet_name='Sheet1')
        
        workbook = writer.book
        worksheet = writer.sheets['Sheet1']
        
        default_font = Font(name='Calibri', size=11)
        default_fill = PatternFill()
        default_border = Border()
        default_alignment = Alignment()
        
        for row in worksheet.iter_rows():
            for cell in row:
                cell.font = default_font
                cell.fill = default_fill
                cell.border = default_border
                cell.alignment = default_alignment
                cell.number_format = '@' 

class DataMergerApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Amazon WHS - Apex Importer")
        self.root.geometry("550x650")

        self.notebook = ttk.Notebook(root)
        self.notebook.pack(fill="both", expand=True, padx=10, pady=10)

        self.tab1 = ttk.Frame(self.notebook)
        self.notebook.add(self.tab1, text="Apex Generator")

        self.tab2 = ttk.Frame(self.notebook)
        self.notebook.add(self.tab2, text="Atualização da Base")

        self.file1_path = tk.StringVar()
        self.file2_path = tk.StringVar()
        self.csv_current_path = tk.StringVar()
        self.xlsx_apex_base_path = tk.StringVar()

        self.setup_tab1()
        self.setup_tab2()
        self.load_logo()

    def load_logo(self):
        try:
            logo_path = r"C:\Users\vitor\Desktop\Apex-generator\image.png"
            if not os.path.exists(logo_path):
                logo_path = "image.png"

            if os.path.exists(logo_path):
                logo_image = Image.open(logo_path)
                if hasattr(Image, 'Resampling'):
                    logo_image = logo_image.resize((200, 80), Image.Resampling.LANCZOS)
                else:
                    logo_image = logo_image.resize((200, 80), Image.ANTIALIAS)
                self.logo = ImageTk.PhotoImage(logo_image)
                tk.Label(self.tab1, image=self.logo).pack(pady=10)
                tk.Label(self.tab2, image=self.logo).pack(pady=10)
            else:
                tk.Label(self.tab1, text="[Amazon WHS]").pack(pady=10)
                tk.Label(self.tab2, text="[Amazon WHS]").pack(pady=10)
        except Exception as e:
            print("Erro logo:", str(e))
            tk.Label(self.tab1, text="[Logo não carregada]").pack()
            tk.Label(self.tab2, text="[Logo não carregada]").pack()

    def setup_tab1(self):
        tk.Label(self.tab1, text="Geração de Arquivos Apex", font=("Arial", 14, "bold")).pack(pady=10)
        tk.Label(self.tab1, text="Arquivo RFID (CSV):").pack()
        tk.Entry(self.tab1, textvariable=self.file1_path, width=60).pack()
        tk.Button(self.tab1, text="Selecionar arquivo", command=self.load_file1).pack(pady=5)

        tk.Label(self.tab1, text="Planilha Associados (XLSX):").pack()
        tk.Entry(self.tab1, textvariable=self.file2_path, width=60).pack()
        tk.Button(self.tab1, text="Selecionar arquivo", command=self.load_file2).pack(pady=5)

        tk.Button(self.tab1, text="Mesclar dados", command=self.process_data, 
                 bg="#0077cc", fg="white", font=("Arial", 10, "bold")).pack(pady=20)

    def setup_tab2(self):
        tk.Label(self.tab2, text="Atualização da Base Apex", font=("Arial", 14, "bold")).pack(pady=10)
        instruction_text = (
            "Esta função atualiza a base do Apex comparando:\n"
            "• Arquivo CSV atual com funcionários ativos\n"
            "• Base de dados atual do Apex\n\n"
            "Resultado: Nova planilha com status atualizados"
        )
        tk.Label(self.tab2, text=instruction_text, justify="left", 
                font=("Arial", 9), wraplength=500).pack(pady=10)

        tk.Label(self.tab2, text="Arquivo CSV Atual (Funcionários Ativos):").pack()
        tk.Entry(self.tab2, textvariable=self.csv_current_path, width=60).pack()
        tk.Button(self.tab2, text="Selecionar CSV", command=self.load_csv_current).pack(pady=5)

        tk.Label(self.tab2, text="Base de Dados Apex (XLSX):").pack()
        tk.Entry(self.tab2, textvariable=self.xlsx_apex_base_path, width=60).pack()
        tk.Button(self.tab2, text="Selecionar Base Apex", command=self.load_xlsx_apex_base).pack(pady=5)

        tk.Button(self.tab2, text="Atualizar Base", command=self.update_apex_base, 
                 bg="#00AA44", fg="white", font=("Arial", 10, "bold")).pack(pady=20)

    def load_file1(self):
        path = filedialog.askopenfilename(filetypes=[["CSV Files", "*.csv"]])
        if path: self.file1_path.set(path)

    def load_file2(self):
        path = filedialog.askopenfilename(filetypes=[["Excel Files", "*.xlsx"]])
        if path: self.file2_path.set(path)

    def load_csv_current(self):
        path = filedialog.askopenfilename(filetypes=[["CSV Files", "*.csv"]])
        if path: self.csv_current_path.set(path)

    def load_xlsx_apex_base(self):
        path = filedialog.askopenfilename(filetypes=[["Excel Files", "*.xlsx"]])
        if path: self.xlsx_apex_base_path.set(path)

    def process_data(self):
        try:
            # Lê o CSV com codificação UTF-8-SIG para evitar erros de charmap
            with open(self.file1_path.get(), 'r', encoding='utf-8-sig') as file:
                content = file.readlines()
                delimiter = ',' if ',' in content[0] else ';'

            df1 = pd.read_csv(self.file1_path.get(), delimiter=delimiter, dtype=str, encoding='utf-8-sig')
            
            if not {'Employee ID', 'Badge ID', 'Login'}.issubset(df1.columns):
                messagebox.showerror("Erro", "Arquivo CSV deve conter 'Employee ID', 'Badge ID' e 'Login'")
                return

            df2 = pd.read_excel(self.file2_path.get(), skiprows=1, dtype=str)
            
            # CORREÇÃO AQUI: Adicionado str(col) para evitar erro com colunas numéricas
            df2 = df2.drop(columns=[col for col in df2.columns if str(col).strip().lower() == 'employee id'], errors='ignore')
            
            if not {'Nome', 'Login'}.issubset(df2.columns):
                messagebox.showerror("Erro", "Arquivo Excel deve conter 'Nome' e 'Login'")
                return

            df2 = df2.dropna(subset=['Nome'])
            df2[['First Name', 'Last Name']] = df2['Nome'].str.split(' ', n=1, expand=True)
            df2['Last Name'].fillna('', inplace=True)

            merged = pd.merge(df2, df1[['Login', 'Employee ID', 'Badge ID']], on='Login', how='left')
            result = merged[['First Name', 'Last Name', 'Employee ID', 'Badge ID']].copy()

            result = result.rename(columns={
                'Employee ID': 'Employee ID*',
                'First Name': 'First Name*',
                'Last Name': 'Last Name*',
                'Badge ID': 'Badge #'
            })

            result['Language'] = 'ENGLISH'
            result['Active*'] = 'Y'
            result['Reporting Group'] = 'Associados'
            result['User Group 2'] = ''
            result['User Group 3'] = ''
            result['User Group 4'] = ''
            result['User Group 5'] = ''
            result['Expiration Date'] = ''

            result = result[
                ['Employee ID*', 'First Name*', 'Last Name*', 'Badge #', 
                'Language', 'Active*', 'Reporting Group', 
                'User Group 2', 'User Group 3', 'User Group 4', 'User Group 5', 'Expiration Date']
            ]

            save_path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx")], initialfile=f"planilha_completa_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
            if not save_path:
                messagebox.showwarning("Aviso", "Arquivo não foi salvo.")
                return

            salvar_excel_sem_formatacao(result, save_path)

            erros = merged[merged['Badge ID'].isna()].copy()
            if not erros.empty:
                erros = erros.rename(columns={
                    'Employee ID': 'Employee ID*',
                    'First Name': 'First Name*',
                    'Last Name': 'Last Name*',
                    'Badge ID': 'Badge #'
                })
                error_path = os.path.splitext(save_path)[0] + "_erros.xlsx"
                salvar_excel_sem_formatacao(erros, error_path)

            pdf_path = os.path.splitext(save_path)[0] + "_termos.pdf"
            gerar_pdf_termos(result, df1, pdf_path)

            instance_df = pd.DataFrame({
                'Instance Name': result['First Name*'] + ' ' + result['Last Name*'],
                'Instance Num': result['Employee ID*'],
                'Desc': ''
            })
            instance_path = os.path.splitext(save_path)[0] + "_Instances.xlsx"
            salvar_excel_sem_formatacao(instance_df, instance_path)

            messagebox.showinfo("Concluído", f"Arquivos salvos com sucesso!")

        except Exception as e:
            messagebox.showerror("Erro", str(e))

    def update_apex_base(self):
        try:
            if not self.csv_current_path.get() or not self.xlsx_apex_base_path.get():
                messagebox.showerror("Erro", "Selecione ambos os arquivos!")
                return

            # CORREÇÃO: Encoding UTF-8-SIG também aqui
            with open(self.csv_current_path.get(), 'r', encoding='utf-8-sig') as file:
                content = file.readlines()
                delimiter = ',' if ',' in content[0] else ';'

            df_csv_current = pd.read_csv(self.csv_current_path.get(), delimiter=delimiter, dtype=str, encoding='utf-8-sig')
            
            if 'Employee ID' not in df_csv_current.columns:
                messagebox.showerror("Erro", "CSV deve conter 'Employee ID'")
                return
                
            if 'First Name' not in df_csv_current.columns or 'Last Name' not in df_csv_current.columns:
                if 'Nome' in df_csv_current.columns:
                    df_csv_current[['First Name', 'Last Name']] = df_csv_current['Nome'].str.split(' ', n=1, expand=True)
                    df_csv_current['Last Name'].fillna('', inplace=True)
                else:
                    messagebox.showerror("Erro", "CSV deve conter 'First Name' e 'Last Name' ou 'Nome'")
                    return

            if 'Badge ID' in df_csv_current.columns:
                df_csv_current = df_csv_current.rename(columns={'Badge ID': 'Badge #'})

            df_apex_base = pd.read_excel(self.xlsx_apex_base_path.get(), dtype=str)
            
            employee_id_col = None
            first_name_col = None
            last_name_col = None
            badge_col = None
            active_col = None

            for col in df_apex_base.columns:
                # CORREÇÃO: Convertido para str para evitar erro com int
                col_str = str(col).lower()
                if 'employee id' in col_str: employee_id_col = col
                elif 'first name' in col_str: first_name_col = col
                elif 'last name' in col_str: last_name_col = col
                elif 'badge' in col_str: badge_col = col
                elif 'active' in col_str: active_col = col

            if not all([employee_id_col, first_name_col, last_name_col, badge_col, active_col]):
                messagebox.showerror("Erro", "Base Apex incompleta.")
                return

            column_mapping = {}
            if employee_id_col != 'Employee ID*': column_mapping[employee_id_col] = 'Employee ID*'
            if first_name_col != 'First Name*': column_mapping[first_name_col] = 'First Name*'
            if last_name_col != 'Last Name*': column_mapping[last_name_col] = 'Last Name*'
            if badge_col != 'Badge #': column_mapping[badge_col] = 'Badge #'
            if active_col != 'Active*': column_mapping[active_col] = 'Active*'

            if column_mapping:
                df_apex_base = df_apex_base.rename(columns=column_mapping)

            result_df = df_apex_base.copy()

            csv_employee_ids = set(df_csv_current['Employee ID'].fillna('').astype(str))
            apex_employee_ids = set(df_apex_base['Employee ID*'].fillna('').astype(str))

            novos_usuarios = csv_employee_ids - apex_employee_ids
            
            if novos_usuarios:
                df_novos = df_csv_current[df_csv_current['Employee ID'].isin(novos_usuarios)].copy()
                df_novos = df_novos.rename(columns={'Employee ID': 'Employee ID*', 'First Name': 'First Name*', 'Last Name': 'Last Name*'})
                df_novos['Language'] = 'ENGLISH'
                df_novos['Active*'] = 'Y'
                df_novos['Reporting Group'] = 'Associados'
                
                for col in result_df.columns:
                    if col not in df_novos.columns: df_novos[col] = ''
                df_novos = df_novos[result_df.columns]
                result_df = pd.concat([result_df, df_novos], ignore_index=True)

            usuarios_desativar = apex_employee_ids - csv_employee_ids
            if usuarios_desativar:
                result_df.loc[result_df['Employee ID*'].isin(usuarios_desativar), 'Active*'] = 'N'

            usuarios_comuns = csv_employee_ids & apex_employee_ids
            reativados_count = 0
            if usuarios_comuns:
                mask_reativar = (result_df['Employee ID*'].isin(usuarios_comuns)) & (result_df['Active*'] == 'N')
                result_df.loc[mask_reativar, 'Active*'] = 'Y'
                reativados_count = mask_reativar.sum()

            result_df = result_df.drop_duplicates(subset=['Employee ID*'], keep='last')

            save_path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx")], initialfile=f"apex_base_atualizada_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
            
            if not save_path: return

            salvar_excel_sem_formatacao(result_df, save_path)

            total_registros = len(result_df)
            novos_count = len(novos_usuarios) if novos_usuarios else 0
            desativados_count = len(usuarios_desativar) if usuarios_desativar else 0

            messagebox.showinfo("Atualização Concluída", 
                f"Resumo:\nTotal: {total_registros}\nNovos: {novos_count}\nDesativados: {desativados_count}\nReativados: {reativados_count}")

        except Exception as e:
            messagebox.showerror("Erro", f"Erro no processamento: {str(e)}")

if __name__ == "__main__":
    root = tk.Tk()
    root.withdraw()
    messagebox.showinfo("Instruções", "Selecione os arquivos conforme solicitado nas abas.")
    root.deiconify()
    app = DataMergerApp(root)
    root.mainloop()
